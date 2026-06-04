import io
from datetime import datetime

from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

from app.models.schemas import JobRecord


def _section_lines(job: JobRecord) -> list[str]:
    lines: list[str] = [
        f"PM Content AI — Export Bundle",
        f"Generated: {datetime.utcnow().strftime('%Y-%m-%d %H:%M UTC')}",
        f"Source file: {job.file.filename}",
        "",
    ]
    if job.analysis:
        lines += [
            "=== CONTENT ANALYSIS ===",
            f"Summary: {job.analysis.summary}",
            f"Topics: {', '.join(job.analysis.topics)}",
            f"Transcript excerpt: {job.analysis.transcript[:1500]}",
            "",
        ]
    if job.classification:
        lines += [
            "=== CLASSIFICATION ===",
            f"Pillar: {job.classification.pillar} ({job.classification.pillar_confidence}%)",
            f"Format: {job.classification.format} ({job.classification.format_confidence}%)",
            "",
        ]
    if job.assets:
        lines += [
            "=== REEL SCRIPT ===",
            f"Hook: {job.assets.reel.hook}",
            f"Body: {job.assets.reel.body}",
            f"CTA: {job.assets.reel.cta}",
            "",
            "=== CAPTION ===",
            f"Hook: {job.assets.caption.hook}",
            f"Story: {job.assets.caption.story}",
            f"Recipe: {job.assets.caption.recipe}",
            f"CTA: {job.assets.caption.cta}",
            f"Hashtags: {' '.join(job.assets.caption.hashtags)}",
            f"Tags: {' '.join(job.assets.caption.tags)}",
            "",
            "=== CAROUSEL ===",
        ]
        c = job.assets.carousel
        for i, slide in enumerate([c.slide_1, c.slide_2, c.slide_3, c.slide_4, c.slide_5], 1):
            lines.append(f"Slide {i}: {slide}")
        lines.append("")
        lines += ["=== STORIES ==="]
        for story in job.assets.stories:
            lines.append(f"- {story.type}: {story.content}")
        lines.append("")
    if job.recommendation:
        r = job.recommendation
        lines += [
            "=== PUBLISHING RECOMMENDATION ===",
            f"Platform: {r.platform}",
            f"Content type: {r.content_type}",
            f"Day: {r.recommended_day} {r.recommended_time}",
            f"Series: {r.series}",
            f"Reason: {r.reason}",
            f"Clearance: {r.clearance_flag}",
        ]
    return lines


def export_docx(job: JobRecord) -> bytes:
    doc = Document()
    doc.add_heading("PM Content AI — Export Bundle", 0)
    doc.add_paragraph(f"Source: {job.file.filename}")
    doc.add_paragraph(f"Generated: {datetime.utcnow().strftime('%Y-%m-%d %H:%M UTC')}")

    for line in _section_lines(job)[3:]:
        if line.startswith("==="):
            doc.add_heading(line.strip("= "), level=2)
        elif line == "":
            continue
        else:
            doc.add_paragraph(line)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def export_pdf(job: JobRecord) -> bytes:
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, rightMargin=inch, leftMargin=inch, topMargin=inch, bottomMargin=inch)
    styles = getSampleStyleSheet()
    body = ParagraphStyle("Body", parent=styles["Normal"], fontSize=10, leading=14, spaceAfter=8)
    story = []
    for line in _section_lines(job):
        if line.startswith("==="):
            story.append(Spacer(1, 0.15 * inch))
            story.append(Paragraph(line.strip("= "), styles["Heading2"]))
        elif line == "":
            story.append(Spacer(1, 0.08 * inch))
        else:
            safe = line.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            story.append(Paragraph(safe, body))
    doc.build(story)
    return buffer.getvalue()
